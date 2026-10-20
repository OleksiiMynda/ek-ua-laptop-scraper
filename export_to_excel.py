import json
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open("laptops_ek_ua.json", "r", encoding="utf-8") as f:
    data = json.load(f)

wb = openpyxl.Workbook()

ws_data = wb.active
ws_data.title = "Дані"

headers = ["Назва", "Опис", "Ціна (грн)", "Місто", "Магазин"]
ws_data.append(headers)

header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
border_thin = Side(border_style="thin", color="D9D9D9")
cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

for col_num in range(1, len(headers) + 1):
    cell = ws_data.cell(row=1, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row_idx, item in enumerate(data, start=2):
    ws_data.append([
        item.get("назва", ""),
        item.get("опис", ""),
        item.get("ціна", 0),
        item.get("місто", ""),
        item.get("магазин", "")
    ])
    for col_num in range(1, 6):
        c = ws_data.cell(row=row_idx, column=col_num)
        c.border = cell_border
        if col_num == 3:
            c.number_format = "#,##0 ₴"
            c.alignment = Alignment(horizontal="right")

for col in ws_data.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_data.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

ws_pivot = wb.create_sheet(title="Зведена таблиця")

ws_pivot.merge_cells("A1:C1")
title_cell = ws_pivot["A1"]
title_cell.value = "Максимальна ціна ноутбуків для магазинів (> 3 пропозицій)"
title_cell.font = Font(name="Calibri", size=13, bold=True, color="1F4E79")
title_cell.alignment = Alignment(horizontal="left", vertical="center")

pivot_headers = ["Назва магазину", "Кількість пропозицій", "Максимальна ціна (грн)"]
ws_pivot.append([])  
ws_pivot.append(pivot_headers)  

pivot_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
for col_num in range(1, 4):
    cell = ws_pivot.cell(row=3, column=col_num)
    cell.font = header_font
    cell.fill = pivot_header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

shop_prices = defaultdict(list)
for item in data:
    shop_prices[item["магазин"]].append(item["ціна"])

filtered_shops = [
    (shop, len(prices), max(prices))
    for shop, prices in shop_prices.items()
    if len(prices) > 3
]
filtered_shops.sort(key=lambda x: x[1], reverse=True)

start_row = 4
for idx, (shop, count, max_p) in enumerate(filtered_shops, start=start_row):
    ws_pivot.append([shop, count, max_p])
    
    c_shop = ws_pivot.cell(row=idx, column=1)
    c_count = ws_pivot.cell(row=idx, column=2)
    c_price = ws_pivot.cell(row=idx, column=3)
    
    for c in (c_shop, c_count, c_price):
        c.border = cell_border
        c.font = Font(name="Calibri", size=11)
        
    c_count.alignment = Alignment(horizontal="center")
    c_price.number_format = "#,##0 ₴"
    c_price.alignment = Alignment(horizontal="right")

total_row = start_row + len(filtered_shops)
ws_pivot.cell(row=total_row, column=1, value="Загальний підсумок").font = Font(name="Calibri", size=11, bold=True)
ws_pivot.cell(row=total_row, column=2, value=f"=SUM(B{start_row}:B{total_row - 1})").font = Font(name="Calibri", size=11, bold=True)
ws_pivot.cell(row=total_row, column=2).alignment = Alignment(horizontal="center")
ws_pivot.cell(row=total_row, column=3, value=f"=MAX(C{start_row}:C{total_row - 1})").font = Font(name="Calibri", size=11, bold=True)
ws_pivot.cell(row=total_row, column=3).number_format = "#,##0 ₴"
ws_pivot.cell(row=total_row, column=3).alignment = Alignment(horizontal="right")

for col_num in range(1, 4):
    c = ws_pivot.cell(row=total_row, column=col_num)
    c.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    c.border = cell_border

ws_pivot.column_dimensions["A"].width = 28
ws_pivot.column_dimensions["B"].width = 24
ws_pivot.column_dimensions["C"].width = 26

output_excel = "laptops_analysis.xlsx"
wb.save(output_excel)
print(f"Файл {output_excel} успішно створено!")
